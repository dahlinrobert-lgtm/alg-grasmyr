#!/usr/bin/env python3
"""Inspect an Älgdata export before writing a source-specific normalizer.

Usage:
    python scripts/inspect_export.py path/to/export.xlsx
    python scripts/inspect_export.py path/to/export.csv

This tool deliberately does not assume a particular schema until a genuine
export has been inspected.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path


def inspect_csv(path: Path) -> None:
    # Swedish Excel exports may use semicolon-delimited CSV and UTF-8 BOM.
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = raw[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"
    rows = list(csv.reader(raw.splitlines(), delimiter=delimiter))
    print(json.dumps({
        "file": str(path),
        "format": "csv",
        "delimiter": repr(delimiter),
        "row_count_including_header": len(rows),
        "header": rows[0] if rows else [],
        "sample_rows": rows[1:6],
    }, ensure_ascii=False, indent=2))


def inspect_excel(path: Path) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("Install dependency first: python -m pip install openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    result = {"file": str(path), "format": "xlsx", "sheets": []}
    for ws in wb.worksheets:
        iterator = ws.iter_rows(values_only=True)
        sample = []
        for _, row in zip(range(6), iterator):
            sample.append(list(row))
        result["sheets"].append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sample_rows": sample,
        })
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path, help="Downloaded CSV or XLSX export")
    args = parser.parse_args()
    path = args.file
    if not path.exists() or not path.is_file():
        parser.error(f"File does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        inspect_csv(path)
    elif suffix in {".xlsx", ".xlsm"}:
        inspect_excel(path)
    elif suffix == ".xls":
        raise SystemExit("Legacy .xls is not supported by this script; export as .xlsx or CSV.")
    else:
        raise SystemExit(f"Unsupported file extension: {suffix}")


if __name__ == "__main__":
    main()
